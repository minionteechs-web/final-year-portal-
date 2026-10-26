# -*- coding: utf-8 -*-
import sys
import re
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
from docx import Document

ROOT = Path(r"c:\Users\USER\Desktop\Final Year Project")
cands = [
    ROOT / "Eze_Chapters_1-3_UPDATED.docx",
    ROOT / "Eze_Chapters_1-3_FIGURES.docx",
    ROOT / "Eze_Dominion_Chibuikem_Chapters_1-3_SYNC.docx",
    ROOT / "Eze_Dominion_Chibuikem_Chapters_1-3.docx",
]
doc_path = next(p for p in cands if p.exists())
d = Document(str(doc_path))
print("Scanning Word:", doc_path.name)

print("\n=== Spaced hyphen ' - ' ===")
found_spaced = False
for i, p in enumerate(d.paragraphs):
    t = p.text
    if " - " in t:
        found_spaced = True
        for m in re.finditer(r".{0,45} - .{0,45}", t):
            print(f"  p{i}: {m.group()!r}")
if not found_spaced:
    print("  None found.")

print("\n=== Em dash (—) or en dash (–) ===")
for i, p in enumerate(d.paragraphs):
    t = p.text
    if "—" in t or "–" in t:
        for m in re.finditer(r".{0,40}[—–].{0,40}", t):
            print(f"  p{i}: {m.group()!r}")

print("\n=== Markdown leftovers (** ` etc.) ===")
for i, p in enumerate(d.paragraphs):
    t = p.text
    if "**" in t or "`" in t or t.strip().startswith("#"):
        print(f"  p{i}: {t[:120]!r}")

print("\n=== Odd replacement chars / mojibake ===")
for i, p in enumerate(d.paragraphs):
    t = p.text
    if re.search(r"Ã.|â€|ï¿½|Â |â€™|â€œ", t):
        print(f"  p{i}: {t[:120]!r}")

print("\n=== Markdown sources: ' - ' lines ===")
for fn in [
    "00-preliminary-pages.md",
    "chapter-1-introduction.md",
    "chapter-2-literature-review.md",
    "chapter-3-research-methodology.md",
    "references.md",
]:
    text = (ROOT / fn).read_text(encoding="utf-8")
    for li, line in enumerate(text.splitlines(), 1):
        if " - " in line and "-----" not in line and not line.strip().startswith("|"):
            print(f"  {fn}:{li}: {line.strip()[:120]}")

print("\n=== Markdown sources: em dashes ===")
for fn in [
    "00-preliminary-pages.md",
    "chapter-1-introduction.md",
    "chapter-2-literature-review.md",
    "chapter-3-research-methodology.md",
]:
    text = (ROOT / fn).read_text(encoding="utf-8")
    for li, line in enumerate(text.splitlines(), 1):
        if "—" in line or "–" in line:
            print(f"  {fn}:{li}: {line.strip()[:120]}")
